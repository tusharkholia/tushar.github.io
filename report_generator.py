"""Report assembly and export utilities for PE Lens."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from data_collection import DealRecord


def _safe_list(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(v) for v in value]
    return []


def render_markdown_report(query: str, analysis: dict, deals: list[DealRecord]) -> str:
    title = analysis.get("title") or f"PE Lens Intelligence Report: {query}"

    lines = [
        f"# TITLE: {title}",
        "",
        "## EXECUTIVE SUMMARY",
    ]
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("executive_summary"))])

    lines.append("\n## MARKET CONTEXT")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("market_context"))])

    lines.append("\n## RECENT DEAL SNAPSHOT")
    lines.append("| Date | Target | Acquirer | Deal Size | Valuation Multiple | Strategic Rationale |")
    lines.append("|---|---|---|---|---|---|")
    for d in deals:
        lines.append(
            f"| {d.date} | {d.target} | {d.acquirer} | {d.deal_size} | {d.valuation_multiple} | {d.strategic_rationale} |"
        )

    lines.append("\n## VALUATION INSIGHTS")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("valuation_insights"))])

    lines.append("\n## BUYER LANDSCAPE")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("buyer_landscape"))])

    lines.append("\n## KEY THEMES")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("key_themes"))])

    lines.append("\n## RISKS & HEADWINDS")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("risks_headwinds"))])

    lines.append("\n## FORWARD OUTLOOK")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("forward_outlook"))])

    lines.append("\n## ANALYST TAKE")
    lines.extend([f"- {item}" for item in _safe_list(analysis.get("analyst_take"))])

    return "\n".join(lines)


def deals_csv_bytes(deals: list[DealRecord]) -> bytes:
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(
        [
            "Date",
            "Target",
            "Acquirer",
            "Deal Size",
            "Valuation Multiple",
            "Strategic Rationale",
            "Revenue",
            "Buyer Type",
            "Sector",
            "Geography",
            "Source",
            "Source URL",
        ]
    )
    for d in deals:
        writer.writerow(
            [
                d.date,
                d.target,
                d.acquirer,
                d.deal_size,
                d.valuation_multiple,
                d.strategic_rationale,
                d.revenue,
                d.buyer_type,
                d.sector,
                d.geography,
                d.source_title,
                d.source_url,
            ]
        )
    return stream.getvalue().encode("utf-8")


def report_to_docx_bytes(report_text: str, title: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for line in report_text.splitlines():
        if line.startswith("## "):
            doc.add_heading(line.replace("## ", ""), level=2)
        elif line.startswith("# "):
            continue
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("|"):
            doc.add_paragraph(line)
        elif line.strip():
            doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def report_to_pdf_bytes(report_text: str, title: str) -> bytes:
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=LETTER)
    width, height = LETTER
    y = height - 42
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(35, y, title[:95])
    y -= 22
    pdf.setFont("Helvetica", 9)

    for raw in report_text.splitlines():
        line = raw.strip()
        if not line:
            y -= 6
            continue
        if y <= 38:
            pdf.showPage()
            y = height - 42
            pdf.setFont("Helvetica", 9)
        for chunk_start in range(0, len(line), 112):
            pdf.drawString(35, y, line[chunk_start : chunk_start + 112])
            y -= 11
            if y <= 38:
                pdf.showPage()
                y = height - 42
                pdf.setFont("Helvetica", 9)

    pdf.save()
    buf.seek(0)
    return buf.read()
