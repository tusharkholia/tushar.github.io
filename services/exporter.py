"""Export helpers for PDF and Word outputs."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


def memo_to_docx_bytes(memo_text: str) -> bytes:
    document = Document()
    document.add_heading("AI Research & Deal Memo", level=1)
    for line in memo_text.splitlines():
        if not line.strip():
            continue
        if line.strip().endswith(":") or line.strip().isupper():
            document.add_heading(line.strip(), level=2)
        else:
            document.add_paragraph(line)

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream.read()


def memo_to_pdf_bytes(memo_text: str) -> bytes:
    stream = BytesIO()
    pdf = canvas.Canvas(stream, pagesize=LETTER)
    width, height = LETTER
    y = height - 50
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(40, y, "AI Research & Deal Memo")
    y -= 24

    pdf.setFont("Helvetica", 10)
    for raw_line in memo_text.splitlines():
        line = raw_line.strip()
        if not line:
            y -= 8
            continue
        for segment in [line[i : i + 105] for i in range(0, len(line), 105)]:
            if y < 40:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 40
            pdf.drawString(40, y, segment)
            y -= 13

    pdf.save()
    stream.seek(0)
    return stream.read()
